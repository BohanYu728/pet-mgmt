"""Build the consolidated report from the user's 00931 template and real evidence."""
from pathlib import Path
from copy import deepcopy
import hashlib, json, re, zipfile
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

ROOT=Path(__file__).resolve().parents[1]
REF=ROOT/'.local/template-review/reference.docx'
SOURCE=ROOT/'docs/宠物管家系统个人技能实训综合报告.md'
OUTPUT=ROOT/'宠物管家系统个人技能实训综合报告_AI图完善版.docx'
QA=ROOT/'.local/template-review'

def source(path): return (ROOT/path).read_text(encoding='utf-8-sig')
def extract(path,start,end):
    s=source(path); a=s.index(start); b=s.index(end,a); return s[a:b].rstrip()

def expand(text):
    codes={
      'admin-create':('java',extract('src/main/java/com/petmgmt/service/AdminUserService.java','    @Transactional\n    public User create','    @Transactional\n    public User update')),
      'role-home':('javascript',extract('frontend/src/roleAccess.js','export function roleHome','export function canVisit')),
      'ci':('yaml',extract('.github/workflows/ci.yml','      - uses: actions/setup-java','      - uses: actions/upload-artifact')),
      'csrf':('java',extract('src/main/java/com/petmgmt/security/SecurityConfig.java','        CookieCsrfTokenRepository csrfRepository','            .cors')),
      'ask':('java',extract('src/main/java/com/petmgmt/ai/rag/AssistantService.java','    public AskResult ask','        try {')),
      'stock':('java',extract('src/main/java/com/petmgmt/service/ShopService.java','            if (productRepository.decrementStockIfAvailable','            OrderItem oi')),
    }
    for name,(language,code) in codes.items():
        block=f'<!-- code:{name}:start -->\n```{language}\n{code}\n```\n<!-- code:{name}:end -->'
        if '@@CODE:'+name+'@@' in text: text=text.replace('@@CODE:'+name+'@@',block)
        else: text=re.sub(r'<!-- code:'+name+r':start -->.*?<!-- code:'+name+r':end -->',lambda _:block,text,flags=re.S)
    shots=json.loads(source('docs/evidence/features/audit.json'))['shots']
    groups={'care':[s for s in shots if s['id'].split('-')[0] in [f'{i:02}' for i in range(1,12)]+['32']],
            'shop':[s for s in shots if s['id'].split('-')[0] in [f'{i:02}' for i in range(12,19)]],
            'management':[s for s in shots if s['id'].split('-')[0] in ['25','26','27','28a','28','29','30','31']],
            'ai':[s for s in shots if s['id'].startswith('ai-')]}
    count=4
    for group,items in groups.items():
        parts=[]
        for i,s in enumerate(items,1):
            count+=1
            path='evidence/features/'+s['id']+'.png'
            if group=='ai': path='evidence/ai-samples/'+s['id'][3:]+'.png'
            if s['id']=='31-billing':path='evidence/demo/billing.png'
            note=s['note']
            if s['id']=='28-merchant-approval':note='当前角色界面显示商户审核列表；既有接口审查另外验证了待审核商户获批后可登录。当前截图不单独证明全部审核过程。'
            section={'care':'5.3','shop':'5.4','management':'5.5','ai':'5.6'}[group]
            parts.append(f'#### {section}.{i} {s["title"]}\n\n入口：{s["route"]}。{note}\n\n![图5-{count} {s["title"]}实际界面]( {path} )')
        block=f'<!-- features:{group}:start -->\n'+ '\n\n'.join(parts)+f'\n<!-- features:{group}:end -->'
        if '@@FEATURES:'+group+'@@' in text: text=text.replace('@@FEATURES:'+group+'@@',block)
        else:text=re.sub(r'<!-- features:'+group+r':start -->.*?<!-- features:'+group+r':end -->',lambda _:block,text,flags=re.S)
    assert count==37
    text=text.replace('图5-40','图5-38').replace('图5-41','图5-39')
    text=text.replace('title、due_date','title、reminder_date').replace('照护提醒及完成状态','照护提醒、已读及忽略状态')
    text=text.replace('提醒的完成状态与记录内容分开维护','提醒的已读、忽略状态与记录内容分开维护')
    return text

def style(doc,name,size,bold=False,outline=None):
    st=doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH) if name not in doc.styles else doc.styles[name]
    st.base_style=doc.styles['Normal'];st.font.name='Times New Roman';st.font.size=Pt(size);st.font.bold=bold
    st.font.color.rgb=RGBColor(0,0,0);st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体')
    f=st.paragraph_format;f.line_spacing=1.5;f.space_after=Pt(6);f.widow_control=True
    if outline is not None:
        f.keep_with_next=True;f.first_line_indent=Cm(0);f.space_before=Pt(12)
        el=OxmlElement('w:outlineLvl');el.set(qn('w:val'),str(outline));st.element.get_or_add_pPr().append(el)
    return st

def main():
    if not REF.exists(): raise SystemExit('先将用户 .doc 模板只读转换为 .local/template-review/reference.docx')
    inv=json.loads((QA/'reference-inventory.json').read_text(encoding='utf-8-sig'))
    assert hashlib.sha256(REF.read_bytes()).hexdigest()==inv['sha256']
    text=expand(SOURCE.read_text(encoding='utf-8-sig'));SOURCE.write_text(text.rstrip()+'\n',encoding='utf8')
    doc=Document(REF);paras=list(doc.paragraphs)
    # Retain the exact cover paragraphs, fonts, page geometry and template package.
    for index,value in [(6,'学号：202432110227'),(7,'姓名：余柏翰'),(8,'班级：____________'),(9,'课程编号：00931')]:
        p=paras[index];props=deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
        p.clear();r=p.add_run(value)
        if props is not None:r._r.insert(0,props)
        p.alignment=1
    p=paras[5];p.add_run('宠物管家系统').font.size=Pt(14);p.alignment=1
    # Empty source content slots can expand; the original remains immutable.
    body=doc._element.body
    for p in paras[11:]: body.remove(p._p)
    style(doc,'报告正文',12).paragraph_format.first_line_indent=Cm(.85)
    style(doc,'报告一级',18,outline=0)
    style(doc,'报告二级',14,bold=True,outline=1)
    style(doc,'报告三级',12,bold=True,outline=2)
    cap=style(doc,'报告图题',10.5);cap.paragraph_format.first_line_indent=Cm(0);cap.paragraph_format.line_spacing=1.15
    code=style(doc,'报告代码',9);code.font.name='Consolas';code.paragraph_format.line_spacing=1.1
    code.paragraph_format.first_line_indent=Cm(0);code.paragraph_format.space_after=Pt(4);code.paragraph_format.alignment=0
    # Template TOC title retained; expandable automatic contents slot.
    p=doc.add_paragraph();fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'TOC \\t "报告一级,1,报告二级,2" \\h \\z')
    r=OxmlElement('w:r');t=OxmlElement('w:t');t.text='目录将在 Word 更新';r.append(t);fld.append(r);p._p.append(fld)
    doc.add_page_break()
    sec=doc.sections[0];sec.different_first_page_header_footer=True
    p=sec.footer.paragraphs[0];p.alignment=1
    p.add_run('— ').font.size=Pt(9);fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');p._p.append(fld);p.add_run(' —').font.size=Pt(9)
    table_count={};images=[];table_rows=[];in_code=False;code_lines=[]
    chapter=0;feature=False
    def paragraph(value,sty='报告正文'):
        p=doc.add_paragraph(value,sty);return p
    def table():
        nonlocal table_rows
        if not table_rows:return
        rows=[r for r in table_rows if not all(re.fullmatch(r':?-+:?',c.strip()) for c in r)]
        table_count[chapter]=table_count.get(chapter,0)+1
        title='实训任务与验收映射' if chapter==1 and table_count[chapter]==1 else {'1':'角色与需求','2':'技术环境','3':'数据库逻辑设计','5':'测试结果与验证范围','6':'AI 工具使用记录'}.get(str(chapter),'项目说明')
        key=rows[0][0].strip()
        if key.startswith('字段（'):title=key[3:-1]+'表关键字段'
        if chapter==5:title={1:'六项 AI 样例的输入与结果',2:'验证记录汇总',3:'代表性测试用例与预期、实际结果'}[table_count[chapter]]
        p=paragraph(f'表{chapter}-{table_count[chapter]} '+title,'报告图题')
        p.alignment=1;p.paragraph_format.keep_with_next=True
        tab=doc.add_table(rows=1,cols=len(rows[0]));tab.autofit=False
        widths=[3.15,5.2,6.3] if len(rows[0])==3 else [2.4,4.25,4.0,4.0]
        for col,w in zip(tab.columns,widths):col.width=Cm(w)
        pr=tab._tbl.tblPr
        borders=OxmlElement('w:tblBorders')
        for edge in ['top','left','bottom','right','insideH','insideV']:
            b=OxmlElement('w:'+edge);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'777777');borders.append(b)
        pr.append(borders)
        for n,row in enumerate(rows):
            cells=tab.rows[0].cells if n==0 else tab.add_row().cells
            for c,v,w in zip(cells,row,widths):
                c.width=Cm(w);p=c.paragraphs[0];p.style=doc.styles['报告图题'];p.add_run('字段' if n==0 and v.strip().startswith('字段（') else v.strip())
                p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1.2
                p.alignment=0
                if len(rows)<=8:p.paragraph_format.keep_with_next=n<len(rows)-1
                elif n==0:p.paragraph_format.keep_with_next=True
                if n==0:
                    for r in p.runs:r.bold=True
                    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EEEEEE');c._tc.get_or_add_tcPr().append(sh)
            trpr=tab.rows[n]._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
            if n==0:trpr.append(OxmlElement('w:tblHeader'))
        paragraph('', '报告图题');table_rows=[]
    for s in text.splitlines():
        if s.startswith('<!--'):continue
        if s.startswith('```'):
            table()
            if in_code:
                p=paragraph('\n'.join(code_lines),'报告代码');p.paragraph_format.keep_together=True
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'F3F3F3');p._p.get_or_add_pPr().append(sh)
                code_lines=[]
            in_code=not in_code;continue
        if in_code:code_lines.append(s);continue
        if s.startswith('|'):
            table_rows.append(s.strip('|').split('|'));continue
        table()
        if not s.strip():continue
        if s.startswith('# '):
            p=paragraph(s[2:],'报告一级');p.alignment=1
            # Project title is source cover-like title, not a TOC chapter.
            for el in list(p._p.get_or_add_pPr().findall(qn('w:outlineLvl'))):p._p.pPr.remove(el)
            p.style=doc.styles['Normal'];p.runs[0].font.size=Pt(18);p.paragraph_format.keep_with_next=True;continue
        if s.startswith('## '):
            chapter=int(s[3]);p=paragraph(s[3:],'报告一级')
            if chapter>1:p.paragraph_format.page_break_before=True
            feature=False;continue
        if s.startswith('### '):
            p=paragraph(s[4:],'报告二级');feature=False
            if re.match(r'### (5\.6|6\.3|6\.5) ',s):p.paragraph_format.page_break_before=True
            continue
        if s.startswith('#### '):
            p=paragraph(s[5:],'报告三级');feature=bool(re.match(r'#### 5\.[3456]\.\d+ ',s))
            if s.startswith('#### 5.9.1 '):p.paragraph_format.page_break_before=True
            if feature and not re.match(r'#### 5\.[345]\.1 ',s):p.paragraph_format.page_break_before=True
            continue
        match=re.match(r'!\[([^]]+)\]\(\s*(.*?)\s*\)',s)
        if match:
            title,path=match.groups();file=ROOT/'docs'/path
            assert file.exists(),file
            w,h=Image.open(file).size
            max_h=15.0 if path.startswith('diagrams/') else (14.7 if 'ai-samples' in path else 12.2)
            width=min(14.65,max_h*w/h)
            p=paragraph('', '报告图题');p.alignment=1;p.paragraph_format.keep_with_next=True
            p.add_run().add_picture(str(file),width=Cm(width))
            p=paragraph(title,'报告图题');p.alignment=1;p.paragraph_format.keep_with_next=False
            images.append({'title':title,'source':str(file.relative_to(ROOT)),'sha256':hashlib.sha256(file.read_bytes()).hexdigest()})
            continue
        p=paragraph(s)
        if feature:p.paragraph_format.keep_with_next=True
        if s.startswith('['):p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.line_spacing=1.2;p.alignment=0
    table()
    doc.core_properties.title='宠物管家系统个人技能实训综合报告';doc.core_properties.author='余柏翰'
    doc.core_properties.subject='按00931课程实验报告模板整合五项实训，采用毕业论文内容深度'
    settings=doc.settings.element
    upd=OxmlElement('w:updateFields');upd.set(qn('w:val'),'true');settings.append(upd)
    doc.save(OUTPUT)
    # python-docx may normalize untouched XML; restore all opaque parts verbatim.
    with zipfile.ZipFile(REF) as a,zipfile.ZipFile(OUTPUT) as b:
        opaque=[n for n in a.namelist() if n.startswith(('customXml/','word/theme/')) or n in ['word/fontTable.xml','word/numbering.xml','word/webSettings.xml']]
        parts={n:b.read(n) for n in b.namelist()}
        for n in opaque:parts[n]=a.read(n)
    with zipfile.ZipFile(OUTPUT,'w',zipfile.ZIP_DEFLATED) as archive:
        for n,data in parts.items():archive.writestr(n,data)
    with zipfile.ZipFile(REF) as a,zipfile.ZipFile(OUTPUT) as b:
        changed=[n for n in a.namelist() if n not in b.namelist() or a.read(n)!=b.read(n)]
        opaque=[n for n in a.namelist() if n.startswith(('customXml/','word/theme/')) or n in ['word/fontTable.xml','word/numbering.xml','word/webSettings.xml']]
        assert all(a.read(n)==b.read(n) for n in opaque),'Unrelated template package changed'
    record={'template_sha256':inv['sha256'],'original_sha256':inv['original_sha256'],'images':images,'uml_diagrams':10,
            'function_showcases':33,'figures':len(images),'tables':sum(table_count.values()),'changed_template_parts':changed,
            'opaque_parts_preserved':True,'report_revision_date':'2026-10-08','test_evidence_date':'2026-09-30'}
    (QA/'build-verification.json').write_text(json.dumps(record,ensure_ascii=False,indent=2),encoding='utf8')
    print('Built',OUTPUT,'figures',len(images),'tables',sum(table_count.values()))

if __name__=='__main__': main()

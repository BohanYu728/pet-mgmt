"""Verify the delivered template report and record artifact-only validation."""
from pathlib import Path
import hashlib, json, re, zipfile
from docx import Document
import pypdfium2 as pdfium

root = Path(__file__).resolve().parents[1]
sha = lambda p: hashlib.sha256(p.read_bytes()).hexdigest()
qa = root / '.local/template-review'
inv = json.loads((qa/'reference-inventory.json').read_text(encoding='utf-8-sig'))
build = json.loads((qa/'build-verification.json').read_text(encoding='utf-8-sig'))
report = root/'宠物管家系统个人技能实训综合报告_AI图完善版.docx'
ref = qa/'reference.docx'
assert sha(ref) == inv['sha256']
assert sha(root/'实验报告参考模板00931.doc') == inv['original_sha256']
doc, template = Document(report), Document(ref)
assert len(doc.sections) == len(template.sections) == 1
for attr in ['page_width','page_height','top_margin','bottom_margin','left_margin','right_margin']:
    assert getattr(doc.sections[0],attr) == getattr(template.sections[0],attr), attr
text = '\n'.join(p.text for p in doc.paragraphs)
assert '@@' not in text and '目录将在 Word 更新' not in text
assert len([p for p in doc.paragraphs if p.style.name == '报告一级']) == 6
assert len(re.findall(r'^5\.[345]\.[0-9]+ ', text, re.M)) == 27
assert len(re.findall(r'^5\.6\.[1-6] ', text, re.M)) == 6
assert all(re.search(r'^4\.'+str(n)+r' ',text,re.M) for n in range(1,6))
assert len(doc.inline_shapes) == 49 and len(doc.tables) == 13
assert all('图3-'+str(n) in text for n in range(7,11))
with zipfile.ZipFile(report) as archive:
    media = {hashlib.sha256(archive.read(n)).hexdigest() for n in archive.namelist() if n.startswith('word/media/')}
    assert all(item['sha256'] in media and sha(root/item['source']) == item['sha256'] for item in build['images'])
    with zipfile.ZipFile(ref) as reference:
        opaque = [n for n in reference.namelist() if n.startswith(('customXml/','word/theme/')) or n in ['word/fontTable.xml','word/numbering.xml','word/webSettings.xml']]
        normalized = [n for n in opaque if archive.read(n) != reference.read(n)]
pdf = pdfium.PdfDocument(root/'.local/report-qa/report.pdf')
pages = len(pdf)
assert pages >= 72
assert all(f'1.{n} ' in text for n in [7,8]) and '5.11 运行结果的阅读方法与问题定位' in text
changed = [n for n in range(1,pages+1) if not (root/f'.local/report-qa-before-code/page-{n}.png').exists() or sha(root/f'.local/report-qa/page-{n}.png') != sha(root/f'.local/report-qa-before-code/page-{n}.png')]
evidence = root/'docs/evidence/features/report-verification.json'
record = json.loads(evidence.read_text(encoding='utf-8-sig'))
record.update(pages=pages, figures=49, uml_diagrams=10, tables=13, sha256=sha(report),
    report_filename=report.name,
    all_images_match=True, template_geometry_matches=True, original_template_unchanged=True,
    chapter_contents_verified=True, five_experiment_structures_verified=True,
    report_revision_date='2026-10-08', test_evidence_date='2026-09-30',
    application_tests_rerun_for_report_revision=False, normalized_by_word=normalized)
evidence.write_text(json.dumps(record,ensure_ascii=False,indent=2)+'\n',encoding='utf8')
print(json.dumps({'pages':pages,'figures':49,'tables':13,'changed_pages':changed,'word_normalized_parts':normalized,'template_unchanged':True},ensure_ascii=False))

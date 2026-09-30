"""Build the single combined report from verified Markdown and real evidence images."""
from pathlib import Path
import argparse
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from PIL import Image
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'docs/宠物管家系统个人技能实训综合报告.md'
OUTPUT = ROOT / '宠物管家系统个人技能实训综合报告.docx'


def font(style, name, size):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '宋体')


def text_runs(p, text):
    # Inline emphasis/code remain editable Word text.
    for i, part in enumerate(re.split(r'(`[^`]+`|\*\*[^*]+\*\*)', text)):
        if part.startswith('**'):
            p.add_run(part[2:-2]).bold = True
        elif part.startswith('`'):
            run = p.add_run(part[1:-1]); run.font.name = 'Consolas'; run.font.size = Pt(9)
        else:
            p.add_run(re.sub(r'\[([^]]+)\]\([^)]+\)', r'\1', part))


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--output', type=Path, default=OUTPUT)
    args = parser.parse_args()
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21); section.page_height = Cm(29.7)
    section.top_margin = Cm(2); section.bottom_margin = Cm(1.8)
    section.left_margin = section.right_margin = Cm(2)
    section.different_first_page_header_footer = True
    section.header_distance = section.footer_distance = Inches(.35)
    for name, size in [('Normal', 10.5), ('Title', 28), ('Subtitle', 13), ('Heading 1', 18), ('Heading 2', 13), ('Heading 3', 11)]:
        font(doc.styles[name], 'Calibri', size)
    for name in ['Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3']:
        doc.styles[name].element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '微软雅黑')
    doc.styles['Subtitle'].font.italic = False
    for style in doc.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    normal = doc.styles['Normal'].paragraph_format
    normal.line_spacing = 1.35; normal.space_after = Pt(7)
    normal.widow_control = True
    for name in ['Heading 1', 'Heading 2', 'Heading 3']:
        f = doc.styles[name].paragraph_format
        f.keep_with_next = True; f.space_before = Pt(12); f.space_after = Pt(7)
    header = section.header.paragraphs[0]
    header.add_run('宠物管家系统  |  个人技能实训').font.size = Pt(9)
    footer = section.footer.paragraphs[0]; footer.alignment = 2
    footer.add_run('第 ')
    fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE'); footer._p.append(fld)
    footer.add_run(' 页')
    lines = SOURCE.read_text(encoding='utf-8-sig').splitlines()
    in_code = False; code = []; table_rows = []; showcase = False; feature_number = 0; compact_summary = False

    def flush_table():
        nonlocal table_rows
        if not table_rows: return
        rows = [r for r in table_rows if not all(re.fullmatch(r':?-+:?', c.strip()) for c in r)]
        table = doc.add_table(rows=1, cols=len(rows[0])); table.style = 'Table Grid'
        table.autofit = False
        widths = [5, 12] if len(rows[0]) == 2 else [4.0, 4.8, 8.2]
        for column, width in zip(table.columns, widths): column.width = Cm(width)
        # Explicit neutral fills/text override theme color.
        for n, row in enumerate(rows):
            cells = table.rows[0].cells if n == 0 else table.add_row().cells
            for cell, value in zip(cells, row):
                text_runs(cell.paragraphs[0], value.strip())
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(5)
                    p.paragraph_format.line_spacing = 1.15
                    for run in p.runs: run.font.size = Pt(9.5); run.font.color.rgb = RGBColor(0,0,0)
                if n == 0:
                    shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'EAF0F3'); cell._tc.get_or_add_tcPr().append(shade)
                    for run in cell.paragraphs[0].runs: run.bold = True
            trpr = table.rows[n]._tr.get_or_add_trPr()
            for cell, width in zip(cells, widths): cell.width = Cm(width)
            trpr.append(OxmlElement('w:cantSplit'))
            if n == 0: trpr.append(OxmlElement('w:tblHeader'))
        doc.add_paragraph().paragraph_format.space_after = Pt(0)
        table_rows = []

    for line in lines:
        if line.startswith('```'):
            flush_table()
            if in_code:
                for index, code_line in enumerate(code):
                    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1
                    p.paragraph_format.left_indent = Cm(.25); p.paragraph_format.right_indent = Cm(.25)
                    p.paragraph_format.keep_with_next = index < len(code) - 1
                    shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'F3F5F7'); p._p.get_or_add_pPr().append(shade)
                    r = p.add_run(code_line); r.font.name = 'Consolas'; r.font.size = Pt(8.5)
                    r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '宋体')
                doc.add_paragraph().paragraph_format.space_after = Pt(0)
                code = []
            in_code = not in_code; continue
        if in_code: code.append(line); continue
        if line.startswith('|'):
            table_rows.append(line.strip('|').split('|')); continue
        flush_table()
        if not line.strip() or line.strip() == '---': continue
        match = re.match(r'!\[([^]]*)\]\(([^)]+)\)', line)
        if match:
            target = SOURCE.parent / match.group(2)
            if not target.exists(): raise FileNotFoundError(target)
            p = doc.add_paragraph(); p.alignment = 1; p.paragraph_format.keep_with_next = True
            with Image.open(target) as image:
                width, height = image.size
            max_height = 3.1 if showcase else (4.5 if height > width else 3.65)
            scale = min(6.65 / width, max_height / height)
            picture = p.add_run().add_picture(str(target), width=Inches(width * scale), height=Inches(height * scale))
            picture._inline.docPr.set('descr', match.group(1))
            caption = doc.add_paragraph(match.group(1)); caption.alignment = 1
            caption.paragraph_format.space_after = Pt(12)
            for r in caption.runs: r.font.size = Pt(9); r.font.color.rgb = RGBColor.from_string('526475')
            continue
        if line.startswith('# '):
            p = doc.add_paragraph(line[2:], 'Title')
            p.paragraph_format.space_before = Pt(78); p.paragraph_format.space_after = Pt(26)
            continue
        if line == '第二部分 个人技能实训':
            p = doc.add_paragraph(line, 'Subtitle'); p.paragraph_format.space_after = Pt(28)
            continue
        if line.startswith('学号 '):
            p = doc.add_paragraph(line); p.paragraph_format.space_after = Pt(40)
            continue
        match = re.match(r'^(#{2,4}) (.*)', line)
        if match:
            level = len(match.group(1)) - 1
            p = doc.add_heading(match.group(2), level)
            if match.group(2) == '综合验证记录': compact_summary = True
            if match.group(2) == '功能逐项验证与展示':
                showcase = True
                compact_summary = False
                p.paragraph_format.page_break_before = True
            if showcase and level == 2 and match.group(2).startswith('功能 '):
                feature_number += 1
                p.paragraph_format.page_break_before = feature_number % 2 == 1
            if level == 1 and (match.group(2).startswith('实验') or match.group(2) in ['项目内容与实现要求', '综合验证记录']):
                p.paragraph_format.page_break_before = True
            continue
        p = doc.add_paragraph()
        if compact_summary:
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(5)
        if line.startswith('- '): p.style = 'List Bullet'; line = line[2:]
        text_runs(p, line)
    flush_table()
    doc.core_properties.title = '宠物管家系统个人技能实训综合报告'
    doc.core_properties.author = '余柏翰'
    doc.core_properties.subject = '第二部分 个人技能实训 实验一至实验五'
    doc.save(args.output)
    print(args.output)

if __name__ == '__main__': main()
